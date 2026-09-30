#!/usr/bin/env python3
"""Generate New Eden top-5 winners deck (revealed 5 → 1)."""
from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

SLUG = os.environ.get("CHALLENGE_SLUG", "new-eden-exchange")
DATABASE_URL = os.environ.get(
    "DATABASE_URL", "postgresql://qtp:qtp@localhost:5432/qtp"
)
OUT = Path(
    os.environ.get(
        "OUT",
        Path(__file__).resolve().parents[1] / "new-eden-exchange-top5-winners.pptx",
    )
)

FONT = "Segoe UI"

# Quantstorm terminal palette
BG = RGBColor(0x1E, 0x1C, 0x1A)
SURFACE = RGBColor(0x2C, 0x29, 0x27)
SURFACE_ALT = RGBColor(0x35, 0x32, 0x2F)
HEADER_BG = RGBColor(0x3E, 0x3A, 0x36)
TEXT = RGBColor(0xF2, 0xEE, 0xE8)
MUTED = RGBColor(0xA8, 0xA0, 0x98)
FAINT = RGBColor(0x78, 0x72, 0x6C)
ACCENT = RGBColor(0xE8, 0x6A, 0x4A)
ACCENT_DIM = RGBColor(0x8C, 0x42, 0x2E)
UP = RGBColor(0x5C, 0xC4, 0xA0)
DOWN = RGBColor(0xE8, 0x6A, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0x4A, 0x46, 0x42)
GOLD = RGBColor(0xE8, 0xC4, 0x6A)
SILVER = RGBColor(0xC8, 0xCC, 0xD0)
BRONZE = RGBColor(0xC9, 0x8A, 0x5C)

RANK_STRIPE = {1: ACCENT, 2: SILVER, 3: BRONZE, 4: ACCENT_DIM, 5: ACCENT_DIM}


def query_final_results() -> tuple[list[dict], list[dict]]:
    sql = f"""
    SELECT final_results
    FROM challenges WHERE slug = '{SLUG}' LIMIT 1;
    """
    raw = subprocess.check_output(
        ["psql", DATABASE_URL, "-t", "-A", "-c", sql],
        text=True,
    ).strip()
    if not raw:
        raise SystemExit(f"No challenge found: {SLUG}")
    rows = sorted(json.loads(raw), key=lambda r: r["rank"])
    top5 = rows[:5]
    if len(top5) < 5:
        raise SystemExit(f"Expected 5 final results, got {len(top5)}")
    return top5, rows


def fmt_money(n: float) -> str:
    sign = "+" if n >= 0 else "-"
    return f"{sign}${abs(n):,.2f}"


def fmt_qty(q: float) -> str:
    if q == int(q):
        return str(int(q))
    return f"{q:,.2f}"


def set_slide_bg(slide, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def style_paragraph(p, *, size: int, bold: bool = False, color: RGBColor = TEXT):
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT


def add_textbox(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 18,
    bold: bool = False,
    color: RGBColor = TEXT,
    align=PP_ALIGN.LEFT,
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    style_paragraph(p, size=size, bold=bold, color=color)
    return box


def add_rect(
    slide,
    left,
    top,
    width,
    height,
    fill: RGBColor,
    *,
    line: RGBColor | None = None,
    radius=None,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    if radius and hasattr(shape, "adjustments") and shape.adjustments:
        shape.adjustments[0] = radius
    return shape


def slide_frame(slide, *, footer: str = "Quantstorm · New Eden Exchange 2026") -> None:
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT)
    add_rect(slide, Inches(0.55), Inches(7.05), Inches(12.2), Inches(0.01), BORDER)
    add_textbox(
        slide,
        Inches(0.55),
        Inches(7.12),
        Inches(6),
        Inches(0.28),
        footer,
        size=10,
        color=FAINT,
    )


def title_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide)
    slide_frame(slide)

    add_textbox(
        slide,
        Inches(0.9),
        Inches(0.85),
        Inches(4),
        Inches(0.35),
        "QUANTSTORM",
        size=11,
        bold=True,
        color=FAINT,
    )
    add_rect(slide, Inches(0.9), Inches(1.55), Inches(2.2), Inches(0.045), ACCENT)
    add_textbox(
        slide,
        Inches(0.9),
        Inches(1.85),
        Inches(11.5),
        Inches(1.1),
        "New Eden Exchange",
        size=52,
        bold=True,
        color=TEXT,
    )
    add_textbox(
        slide,
        Inches(0.9),
        Inches(2.95),
        Inches(11.5),
        Inches(0.7),
        "Top 5 finalists",
        size=30,
        color=ACCENT,
    )
    add_textbox(
        slide,
        Inches(0.9),
        Inches(3.75),
        Inches(9),
        Inches(0.5),
        "Final settlement · Directional scoring · Post-event adjustments applied",
        size=15,
        color=MUTED,
    )

    add_rect(slide, Inches(0.9), Inches(4.55), Inches(11.5), Inches(1.85), SURFACE, line=BORDER)
    add_textbox(
        slide,
        Inches(1.15),
        Inches(4.75),
        Inches(11),
        Inches(0.35),
        "Reveal order",
        size=12,
        bold=True,
        color=ACCENT,
    )
    add_textbox(
        slide,
        Inches(1.15),
        Inches(5.15),
        Inches(11),
        Inches(1.1),
        "#5  →  #4  →  #3  →  #2  →  #1\nFull leaderboard at the end",
        size=20,
        color=TEXT,
    )


def rank_label(rank: int) -> str:
    if rank == 1:
        return "Champion"
    if rank == 2:
        return "Runner-up"
    if rank == 3:
        return "Third place"
    return f"Rank {rank}"


def winner_slide(prs: Presentation, entry: dict) -> None:
    blank = prs.slide_layouts[6]
    rank = entry["rank"]
    name = entry.get("displayName") or entry.get("username", "?")
    username = entry.get("username", "")
    settlement = entry.get("settlement", entry.get("pnl", 0))
    stripe = RANK_STRIPE.get(rank, ACCENT_DIM)

    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide)
    slide_frame(slide, footer=f"Quantstorm · #{rank} {name}")

    add_rect(slide, Inches(0.55), Inches(0.55), Inches(0.1), Inches(6.35), stripe)
    add_rect(slide, Inches(0.75), Inches(0.55), Inches(12.05), Inches(1.35), SURFACE, line=BORDER)

    add_textbox(
        slide,
        Inches(0.95),
        Inches(0.62),
        Inches(1.4),
        Inches(0.45),
        rank_label(rank).upper(),
        size=10,
        bold=True,
        color=stripe,
    )
    add_textbox(
        slide,
        Inches(0.95),
        Inches(0.95),
        Inches(1.5),
        Inches(0.85),
        f"#{rank}",
        size=44,
        bold=True,
        color=TEXT,
    )
    add_textbox(
        slide,
        Inches(2.35),
        Inches(0.88),
        Inches(7.5),
        Inches(0.65),
        name,
        size=34,
        bold=True,
        color=TEXT,
    )
    if username:
        add_textbox(
            slide,
            Inches(2.35),
            Inches(1.48),
            Inches(7.5),
            Inches(0.35),
            f"@{username}",
            size=13,
            color=MUTED,
        )

    add_rect(
        slide, Inches(9.85), Inches(0.82), Inches(2.85), Inches(0.95), HEADER_BG, line=BORDER
    )
    tf = slide.shapes[-1].text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "Settlement"
    p.alignment = PP_ALIGN.CENTER
    style_paragraph(p, size=10, color=MUTED)
    p2 = tf.add_paragraph()
    p2.text = fmt_money(settlement)
    p2.alignment = PP_ALIGN.CENTER
    style_paragraph(
        p2,
        size=22,
        bold=True,
        color=UP if settlement >= 0 else DOWN,
    )

    add_rect(slide, Inches(0.75), Inches(2.05), Inches(5.85), Inches(4.85), SURFACE, line=BORDER)
    add_textbox(
        slide,
        Inches(0.95),
        Inches(2.2),
        Inches(5.4),
        Inches(0.35),
        "Portfolio at close",
        size=13,
        bold=True,
        color=ACCENT,
    )

    cash = entry.get("cash", 0)
    lines = [f"Cash balance    {fmt_money(cash)}"]
    assets = entry.get("assets") or []
    trading = [
        a
        for a in assets
        if a.get("symbol") != "Standard Bond" or abs(a.get("value", 0)) > 1e-6
    ]
    if trading:
        lines.append("")
        for a in trading:
            sym = a["symbol"]
            qty = a["quantity"]
            val = a.get("value", 0)
            mid = a.get("mid", 0)
            side = "long" if qty > 0 else "short" if qty < 0 else "flat"
            lines.append(f"{sym}")
            lines.append(f"  {fmt_qty(qty)} {side}  ·  mid {mid:,.2f}  ·  {fmt_money(val)}")
    else:
        lines.append("")
        lines.append("No open equity positions")
    if any(a.get("symbol") == "Standard Bond" for a in assets):
        lines.append("")
        lines.append("Standard Bond — coupons paid through close")

    symbols = {a["symbol"] for a in trading}
    box = slide.shapes.add_textbox(
        Inches(0.95), Inches(2.6), Inches(5.45), Inches(4.1)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        if not line.strip():
            continue
        if line in symbols:
            style_paragraph(p, size=14, bold=True, color=TEXT)
        elif line.startswith("  "):
            style_paragraph(p, size=12, color=MUTED)
        elif line.startswith("Standard Bond") or line == "No open equity positions":
            style_paragraph(p, size=13, color=MUTED)
        else:
            style_paragraph(p, size=14, bold=True, color=TEXT)

    add_rect(slide, Inches(6.75), Inches(2.05), Inches(6.05), Inches(4.85), SURFACE, line=BORDER)
    add_textbox(
        slide,
        Inches(6.95),
        Inches(2.2),
        Inches(5.6),
        Inches(0.35),
        "Trading activity",
        size=13,
        bold=True,
        color=ACCENT,
    )
    metrics = entry.get("metrics") or {}
    stats = [
        ("Trades executed", f"{metrics.get('trades', 0):,}"),
        ("Volume", f"{metrics.get('volume', 0):,}"),
        (
            "Realized PnL",
            fmt_money(metrics.get("realizedPnl", 0))
            if metrics.get("realizedPnl") is not None
            else "—",
        ),
        ("Spread capture", fmt_money(metrics.get("spreadCapture", 0))),
    ]
    y = 2.75
    for label, value in stats:
        add_rect(slide, Inches(6.95), Inches(y), Inches(5.65), Inches(0.72), SURFACE_ALT)
        add_textbox(
            slide, Inches(7.1), Inches(y + 0.08), Inches(2.5), Inches(0.3), label, size=11, color=MUTED
        )
        add_textbox(
            slide,
            Inches(9.2),
            Inches(y + 0.05),
            Inches(3.2),
            Inches(0.4),
            value,
            size=16,
            bold=True,
            color=TEXT,
            align=PP_ALIGN.RIGHT,
        )
        y += 0.82

    if rank == 1:
        add_rect(slide, Inches(0.75), Inches(6.55), Inches(12.05), Inches(0.42), ACCENT_DIM)
        add_textbox(
            slide,
            Inches(0.75),
            Inches(6.58),
            Inches(12.05),
            Inches(0.38),
            "★  Event champion  ★",
            size=14,
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )


def write_cell(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: RGBColor = WHITE,
    size: int = 11,
    fill: RGBColor = SURFACE,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    tf = cell.text_frame
    tf.margin_left = Pt(8)
    tf.margin_right = Pt(8)
    tf.margin_top = Pt(5)
    tf.margin_bottom = Pt(5)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text = text
    for p in tf.paragraphs:
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.font.name = FONT
        for run in p.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            run.font.name = FONT


def add_leaderboard_slides(prs: Presentation, all_entries: list[dict]) -> None:
    blank = prs.slide_layouts[6]
    page_size = 10
    pages = [
        all_entries[i : i + page_size]
        for i in range(0, len(all_entries), page_size)
    ]
    for page_idx, chunk in enumerate(pages):
        slide = prs.slides.add_slide(blank)
        set_slide_bg(slide)
        slide_frame(slide, footer=f"Quantstorm · Final leaderboard · page {page_idx + 1}")

        title = "Final leaderboard"
        if len(pages) > 1:
            title += f"  ·  {page_idx + 1} / {len(pages)}"
        add_textbox(
            slide,
            Inches(0.75),
            Inches(0.45),
            Inches(10),
            Inches(0.55),
            title,
            size=32,
            bold=True,
            color=TEXT,
        )
        add_rect(slide, Inches(0.75), Inches(1.02), Inches(1.8), Inches(0.04), ACCENT)
        add_textbox(
            slide,
            Inches(0.75),
            Inches(1.12),
            Inches(11),
            Inches(0.3),
            "All traders ranked by final settlement score",
            size=13,
            color=MUTED,
        )

        rows = len(chunk) + 1
        table_shape = slide.shapes.add_table(
            rows, 3, Inches(0.75), Inches(1.5), Inches(11.85), Inches(5.35)
        )
        table = table_shape.table
        tbl = table._tbl
        if tbl.tblPr is not None:
            style_id = getattr(tbl.tblPr, "tableStyleId", None)
            if style_id is not None:
                tbl.tblPr.tableStyleId = None
        table.columns[0].width = Inches(0.75)
        table.columns[1].width = Inches(7.35)
        table.columns[2].width = Inches(3.75)

        for c, label in enumerate(("#", "Trader", "Settlement")):
            write_cell(
                table.cell(0, c),
                label,
                bold=True,
                color=WHITE,
                size=12,
                fill=HEADER_BG,
            )

        for r, entry in enumerate(chunk, start=1):
            rank = entry["rank"]
            name = entry.get("displayName") or entry.get("username", "?")
            user = entry.get("username", "")
            trader = f"{name}  ·  @{user}" if user else name
            settlement = entry.get("settlement", entry.get("pnl", 0))
            row_fill = SURFACE_ALT if r % 2 == 0 else SURFACE
            if rank <= 3:
                row_fill = HEADER_BG
            bold = rank <= 5
            write_cell(
                table.cell(r, 0),
                str(rank),
                bold=bold,
                color=GOLD if rank == 1 else SILVER if rank == 2 else BRONZE if rank == 3 else WHITE,
                size=12,
                fill=row_fill,
            )
            write_cell(
                table.cell(r, 1), trader, bold=bold, color=WHITE, size=12, fill=row_fill
            )
            write_cell(
                table.cell(r, 2),
                fmt_money(settlement),
                bold=bold,
                color=UP if settlement >= 0 else DOWN,
                size=12,
                fill=row_fill,
            )


def build_deck(top5: list[dict]) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    title_slide(prs)
    for entry in sorted(top5, key=lambda e: -e["rank"]):
        winner_slide(prs, entry)
    return prs


def main() -> None:
    top5, all_entries = query_final_results()
    prs = build_deck(top5)
    add_leaderboard_slides(prs, all_entries)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(str(OUT.resolve()))


if __name__ == "__main__":
    main()
